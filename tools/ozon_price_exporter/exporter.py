from __future__ import annotations

import json
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


PRICE_COLUMNS: tuple[tuple[str, str], ...] = (
    ("collected_at", "Дата и время сбора"),
    ("profile_name", "Профиль магазина"),
    ("offer_id", "Артикул продавца"),
    ("product_id", "Product ID Ozon"),
    ("sku", "SKU Ozon"),
    ("name", "Название"),
    ("status", "Статус товара"),
    ("category", "Категория"),
    ("price", "Базовая цена продавца"),
    ("old_price", "Зачёркнутая цена"),
    ("min_price", "Минимальная цена"),
    ("promo_effective_price", "Цена с учётом акций"),
    ("customer_price", "Цена клиента"),
    ("customer_price_source", "Источник цены клиента"),
    ("ozon_card_price", "Цена с Ozon Картой"),
    ("fbo_customer_price", "Цена клиента FBO"),
    ("fbs_customer_price", "Цена клиента FBS"),
    ("currency", "Валюта"),
    ("actions_count", "Количество акций"),
    ("actions_titles", "Акции"),
    ("actions_dates", "Даты окончания акций"),
    ("fbo_stock", "Остаток FBO"),
    ("fbs_stock", "Остаток FBS"),
    ("collection_status", "Статус сбора"),
    ("error", "Ошибка"),
    ("product_url", "Ссылка на страницу цен"),
)

_MONEY_KEYS = {
    "price",
    "old_price",
    "min_price",
    "promo_effective_price",
    "customer_price",
    "ozon_card_price",
    "fbo_customer_price",
    "fbs_customer_price",
}
_INTEGER_KEYS = {"actions_count", "fbo_stock", "fbs_stock"}
_DANGEROUS_PREFIXES = ("=", "+", "-", "@")


def safe_excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (int, float, bool, datetime)):
        return value
    if isinstance(value, (list, tuple, dict)):
        value = json.dumps(value, ensure_ascii=False)
    text = str(value)
    if text.startswith(_DANGEROUS_PREFIXES):
        return "'" + text
    return text


def _apply_sheet_style(sheet, *, table_name: str | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="17324D")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F0")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False

    if table_name and sheet.max_row >= 2 and sheet.max_column >= 1:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table.tableStyleInfo = style
        sheet.add_table(table)

    width_caps = {
        "Название": 42,
        "Категория": 34,
        "Акции": 40,
        "Ошибка": 48,
        "Ссылка на страницу цен": 48,
        "Источник цены клиента": 38,
    }
    for column_index, column_cells in enumerate(sheet.columns, 1):
        header = str(column_cells[0].value or "")
        max_len = len(header)
        for cell in column_cells[1: min(sheet.max_row, 300)]:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 80))
            cell.alignment = Alignment(vertical="top", wrap_text=header in width_caps)
        cap = width_caps.get(header, 24)
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_len + 2, 11), cap)
    sheet.row_dimensions[1].height = 34


def build_workbook_bytes(rows: Iterable[dict[str, Any]], metadata: dict[str, Any]) -> bytes:
    records = list(rows)
    workbook = Workbook()
    prices = workbook.active
    prices.title = "Цены"

    headers = [header for _, header in PRICE_COLUMNS]
    prices.append(headers)
    for record in records:
        prices.append([safe_excel_value(record.get(key)) for key, _ in PRICE_COLUMNS])

    header_to_index = {header: index + 1 for index, (_, header) in enumerate(PRICE_COLUMNS)}
    key_to_index = {key: index + 1 for index, (key, _) in enumerate(PRICE_COLUMNS)}
    for key in _MONEY_KEYS:
        column_index = key_to_index[key]
        for row_index in range(2, prices.max_row + 1):
            prices.cell(row=row_index, column=column_index).number_format = '#,##0.00 "₽"'
    for key in _INTEGER_KEYS:
        column_index = key_to_index[key]
        for row_index in range(2, prices.max_row + 1):
            prices.cell(row=row_index, column=column_index).number_format = "0"
    link_column = header_to_index["Ссылка на страницу цен"]
    for row_index in range(2, prices.max_row + 1):
        cell = prices.cell(row=row_index, column=link_column)
        if isinstance(cell.value, str) and cell.value.startswith("https://"):
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"
    _apply_sheet_style(prices, table_name="OzonCustomerPrices")

    summary = workbook.create_sheet("Сводка")
    summary.append(["Параметр", "Значение"])
    summary_rows = [
        ("Версия приложения", metadata.get("version", "0.1.2")),
        ("Профиль магазина", metadata.get("profile_name")),
        ("Начало сбора", metadata.get("started_at")),
        ("Окончание сбора", metadata.get("completed_at")),
        ("Обнаружено товаров", metadata.get("discovered", len(records))),
        ("Обработано товаров", metadata.get("processed", len(records))),
        ("Успешно", metadata.get("ok", 0)),
        ("С ошибками", metadata.get("failed", 0)),
        ("Остановлено пользователем", bool(metadata.get("stopped"))),
        ("Режим", "Только чтение: приложение не изменяет цены и товары в Ozon"),
        (
            "Цена клиента",
            "Приоритет: marketing_oa_price → marketing_price → встроенный state → отображаемая цена FBO/FBS",
        ),
        (
            "Авторизация",
            "Логин и пароль вводятся только на странице Ozon; приложение сохраняет лишь браузерную сессию профиля",
        ),
    ]
    for item in summary_rows:
        summary.append([safe_excel_value(item[0]), safe_excel_value(item[1])])
    _apply_sheet_style(summary, table_name="OzonPriceExportSummary")
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 88
    for row in summary.iter_rows(min_row=2, max_col=2):
        row[1].alignment = Alignment(vertical="top", wrap_text=True)

    errors = workbook.create_sheet("Ошибки")
    errors.append(["Артикул продавца", "Product ID Ozon", "SKU Ozon", "Ошибка", "Ссылка"])
    for record in records:
        if record.get("collection_status") == "OK" and not record.get("error"):
            continue
        errors.append(
            [
                safe_excel_value(record.get("offer_id")),
                safe_excel_value(record.get("product_id")),
                safe_excel_value(record.get("sku")),
                safe_excel_value(record.get("error") or "Неполные данные"),
                safe_excel_value(record.get("product_url")),
            ]
        )
    _apply_sheet_style(errors, table_name="OzonPriceExportErrors" if errors.max_row >= 2 else None)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def write_workbook(path: Path, rows: Iterable[dict[str, Any]], metadata: dict[str, Any]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = build_workbook_bytes(rows, metadata)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_bytes(payload)
    temporary.replace(path)
    return path
