from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from tools.ozon_price_exporter.browser_process import build_browser_command
from tools.ozon_price_exporter.exporter import build_workbook_bytes, safe_excel_value
from tools.ozon_price_exporter.parser import (
    choose_customer_price,
    extract_common_price_fields,
    extract_products_from_payload,
    extract_total_hint,
    parse_html,
)
from tools.ozon_price_exporter.storage import ProfileStore, profile_slug


def _fixture_html() -> str:
    state = r'''
    {
      "screen": {
        "products": [
          {
            "offerId": "770105",
            "sku": 123456789,
            "title": "Тестовое средство",
            "descriptionCategoryName3Lvl": "Бытовая химия",
            "status": {"stateName": "Продаётся"},
            "price": {
              "price": 338,
              "oldPrice": 483,
              "minAutoPrice": 273,
              "netPrice": 118,
              "marketingSellerPrice": 338,
              "marketingPrice": 299,
              "currency": "RUB"
            },
            "stock": {"present": 69, "reserved": 2},
            "sellerStock": {"present": 99, "reserved": 1}
          }
        ]
      }
    }
    '''
    return f"""
    <html><body>
      <div>249 ₽</div><div>Цена с картой</div>
      <div>299 ₽</div><div>Цена по FBO</div>
      <div>309 ₽</div><div>Цена по FBS</div>
      <script>(window.__MODULE_STATE__=window.__MODULE_STATE__||{{}})["prices"]={state};</script>
    </body></html>
    """


def test_parse_html_extracts_canonical_and_dom_prices() -> None:
    row = parse_html(_fixture_html(), want_offer_id="770105")
    assert row["offer_id"] == "770105"
    assert row["sku"] == "123456789"
    assert row["price"] == 338.0
    assert row["old_price"] == 483.0
    assert row["min_price"] == 273.0
    assert row["marketing_price"] == 299.0
    assert row["ozon_card_price"] == 249.0
    assert row["fbo_customer_price"] == 299.0
    assert row["fbs_customer_price"] == 309.0
    assert row["_parse_error"] is None


def test_extract_products_from_nested_payload_and_total() -> None:
    payload = {
        "result": {
            "totalCount": 2,
            "items": [
                {
                    "item_id": 111,
                    "offer_id": "A-1",
                    "sku": 9001,
                    "marketing_oa_price": "129.90",
                },
                {
                    "productId": "222",
                    "offerId": "A-2",
                    "ozonSku": "9002",
                    "price": {"price": 250, "oldPrice": 300},
                },
            ],
        }
    }
    products = extract_products_from_payload(payload)
    assert set(products) == {"111", "222"}
    assert products["111"]["offer_id"] == "A-1"
    assert products["111"]["customer_price"] == 129.9
    assert products["222"]["price"] == 250.0
    assert extract_total_hint(payload) == 2


def test_current_product_list_payload_is_normalized() -> None:
    payload = {
        "products": [
            {
                "item_id": "2096971288",
                "company_id": "275649",
                "part_item": {
                    "offer_id": "208250 Bug-Off 500мл",
                    "name": "Очиститель следов насекомых для авто 500 мл",
                },
                "part_status": {"texts": {"state_name": "Готов к продаже"}},
                "part_category": {"texts": {"category_name": "Очистители автомобильные"}},
                "part_price": {
                    "price": {"currencyCode": "RUB", "units": "247", "nanos": 0},
                    "old_price": {"currencyCode": "RUB", "units": "353", "nanos": 0},
                    "min_price": {"currencyCode": "RUB", "units": "247", "nanos": 0},
                    "net_price": {"currencyCode": "RUB", "units": "51", "nanos": 0},
                    "currency": "RUB",
                },
                "part_marketing_price": {
                    "price": {"currencyCode": "RUB", "units": "242", "nanos": 0},
                    "seller_price": {"currencyCode": "RUB", "units": "247", "nanos": 0},
                    "oa_price": {"currencyCode": "RUB", "units": "217", "nanos": 0},
                },
                "part_price_indexes_full": {"price_indexes": [{"sku": "2383195686"}]},
                "part_stocks": {
                    "stocks": [
                        {"sku": "2383195686", "source": "fbo", "present": 3, "reserved": 1},
                        {"sku": "2383195686", "source": "fbs", "present": 7, "reserved": 2},
                    ]
                },
            }
        ],
        "total_items": 62,
        "total_count": 55,
        "cursor": "next",
    }

    products = extract_products_from_payload(payload)
    assert set(products) == {"2096971288"}
    row = products["2096971288"]
    assert row["offer_id"] == "208250 Bug-Off 500мл"
    assert row["sku"] == "2383195686"
    assert row["name"] == "Очиститель следов насекомых для авто 500 мл"
    assert row["status"] == "Готов к продаже"
    assert row["category"] == "Очистители автомобильные"
    assert row["price"] == 247.0
    assert row["old_price"] == 353.0
    assert row["min_price"] == 247.0
    assert row["promo_effective_price"] == 242.0
    assert row["customer_price"] == 217.0
    assert row["customer_price_source"] == "list-by-filter.part_marketing_price.oa_price"
    assert row["ozon_card_price"] == 217.0
    assert row["fbo_stock"] == 3
    assert row["fbs_stock"] == 7
    assert extract_total_hint(payload) == 62


def test_common_price_and_precedence() -> None:
    common = extract_common_price_fields(
        {
            "marketing_price": "210",
            "marketing_oa_price": "199",
            "marketing_actions": [
                {"name": "Акция продавца", "date_to": "2026-10-01T00:00:00Z", "is_seller_action": True}
            ],
        }
    )
    assert common["customer_price"] == 199.0
    assert common["promo_effective_price"] == 210.0
    assert common["actions_count"] == 1
    price, source = choose_customer_price({**common, "marketing_price": 220})
    assert price == 199.0
    assert source == "get-common-prices.marketing_oa_price"


def test_profile_isolated_and_path_safe(tmp_path: Path) -> None:
    store = ProfileStore(tmp_path)
    first = store.ensure_profile("Новый магазин")
    second = store.ensure_profile("Другой магазин")
    assert first["slug"] != second["slug"]
    assert Path(first["browser_dir"]).is_dir()
    assert Path(second["browser_dir"]).is_dir()
    assert profile_slug("../../escape").startswith("escape-")
    assert Path(first["browser_dir"]).resolve().is_relative_to((tmp_path / "profiles").resolve())


def test_excel_contract_and_formula_injection_protection() -> None:
    rows = [
        {
            "collected_at": "2026-09-25T10:00:00+02:00",
            "profile_name": "Магазин",
            "offer_id": "=HYPERLINK(\"bad\")",
            "product_id": "111",
            "sku": "9001",
            "name": "Средство",
            "price": 250,
            "customer_price": 199,
            "customer_price_source": "get-common-prices.marketing_oa_price",
            "collection_status": "OK",
            "product_url": "https://seller.ozon.ru/app/prices/manager/111/prices",
        }
    ]
    payload = build_workbook_bytes(
        rows,
        {
            "profile_name": "Магазин",
            "started_at": "2026-09-25T10:00:00+02:00",
            "completed_at": "2026-09-25T10:01:00+02:00",
            "discovered": 1,
            "processed": 1,
            "ok": 1,
            "failed": 0,
        },
    )
    workbook = load_workbook(BytesIO(payload), data_only=False)
    assert workbook.sheetnames == ["Цены", "Сводка", "Ошибки"]
    prices = workbook["Цены"]
    headers = [cell.value for cell in prices[1]]
    assert "Цена клиента" in headers
    offer_column = headers.index("Артикул продавца") + 1
    assert prices.cell(row=2, column=offer_column).value.startswith("'=")
    assert safe_excel_value("@SUM(A1:A2)").startswith("'@")


def test_browser_command_uses_loopback_cdp_without_unsafe_flags(tmp_path: Path) -> None:
    executable = tmp_path / "chrome.exe"
    profile = tmp_path / "profile"
    command = build_browser_command(executable, profile, 9333)

    assert command[0] == str(executable)
    assert "--remote-debugging-port=9333" in command
    assert "--remote-debugging-address=127.0.0.1" in command
    assert f"--user-data-dir={profile}" in command
    assert "--no-sandbox" not in command
    assert "--enable-automation" not in command
